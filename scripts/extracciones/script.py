from __future__ import annotations

import argparse
import csv
import re
from pathlib import Path

import zipfile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException


def normalize_text(value: str) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip().lower()


KEYWORDS = {
    normalize_text("PAGO PSE TU COMPRA SAS"),
    normalize_text("PAGO PSE A Toda Hora  SA"),
}

CANONICAL_DESCRIPTIONS = {
    normalize_text("PAGO PSE TU COMPRA SAS"): "PAGO PSE TU COMPRA SAS",
    normalize_text("PAGO PSE A Toda Hora  SA"): "PAGO PSE A Toda Hora  SA",
}


def clean_currency(value):
    if value is None:
        return 0
    text = str(value).strip()
    if not text:
        return 0
    text = text.replace("$", "").replace(",", "")
    number = float(text)
    if number.is_integer():
        return int(number)
    return number


def extract_matches(input_dir: Path) -> list[dict]:
    input_dir = Path(input_dir)
    results: list[dict] = []

    for workbook_path in sorted(input_dir.glob("*.xlsx")):
        try:
            wb = load_workbook(workbook_path, data_only=True)
        except (InvalidFileException, OSError, zipfile.BadZipFile):
            print(f"Omitiendo archivo no válido: {workbook_path.name}")
            continue

        ws = wb.active
        movimientos_row = None
        for idx, row in enumerate(ws.iter_rows(min_row=1, values_only=True), start=1):
            if isinstance(row[0], str) and row[0].strip() == "Movimientos:":
                movimientos_row = idx
                break

        if movimientos_row is None:
            print(f"No se encontró la sección Movimientos en {workbook_path.name}")
            continue

        header_found = False
        header_row = None
        for idx, row in enumerate(ws.iter_rows(min_row=movimientos_row + 1, values_only=True), start=movimientos_row + 1):
            if row[:6] == ("FECHA", "DESCRIPCIÓN", "SUCURSAL", "DCTO.", "VALOR", "SALDO"):
                header_row = idx
                header_found = True
                break

        if not header_found:
            print(f"No se encontró encabezado de movimientos en {workbook_path.name}")
            continue

        for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
            if not row:
                continue
            fecha = row[0]
            descripcion = row[1]
            valor = row[4]
            if not isinstance(descripcion, str):
                continue

            normalized = normalize_text(descripcion)
            if normalized in KEYWORDS:
                canonical = CANONICAL_DESCRIPTIONS[normalized]
                results.append({
                    "ARCHIVO": workbook_path.name,
                    "FECHA": fecha,
                    "DESCRIPCIÓN": canonical,
                    "VALOR": clean_currency(valor),
                })

    return results


def write_csv(results: list[dict], output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    fieldnames = ["ARCHIVO", "FECHA", "DESCRIPCIÓN", "VALOR"]
    with output_path.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=fieldnames)
        writer.writeheader()
        for row in results:
            writer.writerow({
                "ARCHIVO": row["ARCHIVO"],
                "FECHA": row["FECHA"],
                "DESCRIPCIÓN": row["DESCRIPCIÓN"],
                "VALOR": row["VALOR"],
            })


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Busca ingresos/pagos en Excel y exporta FECHA, DESCRIPCIÓN y VALOR."
    )
    parser.add_argument(
        "--input-dir",
        type=Path,
        default=Path(__file__).resolve().parent / "temp",
        help="Directorio con archivos Excel (.xlsx)."
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=Path(__file__).resolve().parent / "extracciones.csv",
        help="Archivo CSV de salida."
    )
    args = parser.parse_args()

    results = extract_matches(args.input_dir)
    write_csv(results, args.output)

    print(f"Encontradas {len(results)} coincidencias.")
    print(f"Resultado guardado en: {args.output}")

    for row in results:
        print(f"{row['ARCHIVO']} | {row['FECHA']} | {row['DESCRIPCIÓN']} | {row['VALOR']}")


if __name__ == "__main__":
    import zipfile

    main()
